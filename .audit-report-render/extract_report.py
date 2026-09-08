from docx import Document

path = r"C:\Users\longi_vnjgmvk\Downloads\Fiches de non conformité.docx"
document = Document(path)

for paragraph in document.paragraphs:
    if paragraph.text.strip():
        print(paragraph.text)

for table in document.tables:
    print("\n[TABLE]")
    for row in table.rows:
        print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
